"""Owner-scoped, local attachment reading with temporary source retention."""

from __future__ import annotations

import json
import multiprocessing
import re
import shutil
import threading
import time
import uuid
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from queue import Empty
from typing import Any

from .storage import atomic_json, digest, now, search_terms

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_OCR_PAGES = 10
MAX_TEXT_CHARS = 20_000
MAX_DOCX_MEMBERS = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
PENDING_MINUTES = 30
TASK_ID_RE = re.compile(r"[a-f0-9]{12}")
RECORD_ID_RE = re.compile(r"[a-f0-9]{12}")
PAGE_SPEC_RE = re.compile(r"[0-9,，\-—\s]+")
SUPPORTED_SUFFIXES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
}


class AttachmentError(ValueError):
    """A user-safe attachment processing error."""


def format_bytes(size: int) -> str:
    """Render a bounded byte count for a confirmation message."""
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / (1024 * 1024):.1f} MB"


def page_label(pages: list[int]) -> str:
    """Compress page numbers for a compact user-facing message."""
    if not pages:
        return "无"
    ranges = []
    start = previous = pages[0]
    for page in pages[1:]:
        if page == previous + 1:
            previous = page
            continue
        ranges.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = page
    ranges.append(str(start) if start == previous else f"{start}-{previous}")
    return "、".join(ranges)


def parse_page_spec(value: str, *, required: bool = False) -> list[int] | None:
    """Validate a 1-based PDF page selection without accepting free-form text."""
    value = value.strip()
    if not value:
        if required:
            raise AttachmentError("请提供页码，例如：1-3,5。")
        return None
    if not PAGE_SPEC_RE.fullmatch(value):
        raise AttachmentError("页码只支持数字、逗号和范围，例如：1-3,5。")
    pages: set[int] = set()
    for part in re.split(r"[,，]", value):
        part = part.strip().replace("—", "-")
        if not part:
            continue
        if "-" in part:
            bounds = [item.strip() for item in part.split("-")]
            if len(bounds) != 2 or not all(item.isdigit() for item in bounds):
                raise AttachmentError("页码范围格式无效，例如：1-3。")
            start, end = (int(item) for item in bounds)
            if start < 1 or end < start:
                raise AttachmentError("页码范围无效。")
            pages.update(range(start, end + 1))
        elif part.isdigit() and int(part) >= 1:
            pages.add(int(part))
        else:
            raise AttachmentError("页码必须从 1 开始。")
    return sorted(pages)


def _read_text(path: str) -> str:
    data = Path(path).read_bytes()
    if b"\x00" in data[:4096]:
        raise AttachmentError("文本文件编码不受支持，请转换为 UTF-8 后重新发送。")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AttachmentError("文本文件不是 UTF-8 编码，请转换后重新发送。") from exc


def _ensure_text_limit(text: str) -> str:
    if len(text) > MAX_TEXT_CHARS:
        raise AttachmentError(
            "可读取正文超过 20000 字。请发送较小文件；PDF 也可在确认时指定页码。"
        )
    return text.strip()


def _safe_docx(path: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            names = {item.filename for item in infos}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise AttachmentError("文件扩展名为 DOCX，但实际不是有效 Word 文档。")
            if len(infos) > MAX_DOCX_MEMBERS:
                raise AttachmentError("DOCX 内部文件过多，无法安全读取。")
            if any(info.flag_bits & 0x1 for info in infos):
                raise AttachmentError("加密 DOCX 暂不支持，请移除密码后重新发送。")
            if sum(info.file_size for info in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise AttachmentError("DOCX 解压后过大，无法安全读取。")
    except zipfile.BadZipFile as exc:
        raise AttachmentError("文件扩展名为 DOCX，但实际不是有效 Word 文档。") from exc


def detect_file_type(path: str, file_name: str) -> str:
    """Verify both suffix and file signature before a pending task is created."""
    suffix = Path(file_name).suffix.lower()
    if suffix == ".doc":
        raise AttachmentError("旧版 .doc 暂不支持，请另存为 .docx 或 PDF 后发送。")
    file_type = SUPPORTED_SUFFIXES.get(suffix)
    if file_type is None:
        raise AttachmentError("仅支持 PDF、DOCX、HTML 和 Markdown 附件。")
    with Path(path).open("rb") as stream:
        sample = stream.read(8192)
    if file_type == "pdf":
        if not sample.startswith(b"%PDF-"):
            raise AttachmentError("文件扩展名与实际格式不一致，未读取。")
    elif file_type == "docx":
        if not sample.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            raise AttachmentError("文件扩展名与实际格式不一致，未读取。")
        _safe_docx(path)
    else:
        _read_text(path)
    return file_type


def _parse_html(path: str) -> str:
    from bs4 import BeautifulSoup, NavigableString

    soup = BeautifulSoup(_read_text(path), "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    for table in soup.find_all("table"):
        rows = []
        for row in table.find_all("tr"):
            cells = [
                cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])
            ]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            table.replace_with(NavigableString("\n[表格]\n" + "\n".join(rows) + "\n"))
        else:
            table.decompose()
    return _ensure_text_limit(soup.get_text("\n", strip=True))


def _parse_docx(path: str) -> str:
    from docx import Document

    _safe_docx(path)
    try:
        document = Document(path)
    except Exception as exc:  # python-docx has many backend-specific errors.
        raise AttachmentError("DOCX 无法读取，文件可能损坏或受保护。") from exc
    parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _ensure_text_limit("\n".join(parts))


def _ocr_pdf_page(pdf_path: str, page_number: int, model_root: str) -> str:
    """Run local CPU OCR for one rendered page without writing an image file."""
    try:
        import pypdfium2 as pdfium
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise AttachmentError(
            "本地 OCR 组件尚未安装完成，请联系管理员完成部署后重试。"
        ) from exc
    document = pdfium.PdfDocument(pdf_path)
    try:
        page = document[page_number - 1]
        bitmap = page.render(scale=2.0)
        image = bitmap.to_pil()
    finally:
        document.close()
    engine = RapidOCR(
        params={
            "Global.model_root_dir": model_root,
            "Global.log_level": "error",
            "EngineConfig.onnxruntime.use_cuda": False,
        }
    )
    result = engine(image)
    texts = getattr(result, "txts", None) or ()
    return "\n".join(str(text).strip() for text in texts if str(text).strip())


def warm_ocr_models(model_root: str) -> None:
    """Download and validate the local CPU OCR models before the shared host starts."""
    try:
        import numpy
        import pypdfium2  # noqa: F401 - deployment check for PDF rendering.
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise AttachmentError("本地 OCR 依赖安装不完整。") from exc
    engine = RapidOCR(
        params={
            "Global.model_root_dir": model_root,
            "Global.log_level": "error",
            "EngineConfig.onnxruntime.use_cuda": False,
        }
    )
    engine(numpy.full((32, 32, 3), 255, dtype=numpy.uint8))


def _parse_pdf(
    path: str,
    selected_pages: list[int] | None,
    ocr_pages: list[int],
    model_root: str,
) -> dict[str, Any]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
    except Exception as exc:
        raise AttachmentError("PDF 无法读取，文件可能损坏。") from exc
    if reader.is_encrypted:
        raise AttachmentError("加密 PDF 暂不支持，请移除密码后重新发送。")
    total_pages = len(reader.pages)
    if total_pages < 1:
        raise AttachmentError("PDF 没有可读取页面。")
    pages = selected_pages or list(range(1, total_pages + 1))
    if len(pages) > MAX_PDF_PAGES:
        raise AttachmentError("PDF 单次最多读取 100 页，请在确认时指定页码。")
    if any(page > total_pages for page in pages):
        raise AttachmentError(f"页码超出范围；该 PDF 共 {total_pages} 页。")
    if any(page not in pages for page in ocr_pages):
        raise AttachmentError("OCR 页码必须属于本次已选择的 PDF 页码。")
    if len(ocr_pages) > MAX_OCR_PAGES:
        raise AttachmentError("每次最多识别 10 页扫描件，请分批确认 OCR。")
    parts = []
    scanned_pages = []
    for page_number in pages:
        try:
            text = (reader.pages[page_number - 1].extract_text() or "").strip()
        except Exception as exc:
            raise AttachmentError(f"第 {page_number} 页文字提取失败。") from exc
        if len(text) < 20:
            scanned_pages.append(page_number)
            if page_number in ocr_pages:
                text = _ocr_pdf_page(path, page_number, model_root)
                if not text:
                    text = "[该页 OCR 未识别出可用文字]"
            else:
                text = "[该页为扫描页，尚未进行 OCR]"
        parts.append(f"[PDF 第 {page_number} 页]\n{text}")
    return {
        "text": _ensure_text_limit("\n\n".join(parts)),
        "read_pages": pages,
        "scanned_pages": scanned_pages,
        "ocr_pages": ocr_pages,
        "total_pages": total_pages,
    }


def _parse_worker(payload: dict[str, Any]) -> dict[str, Any]:
    file_type = payload["file_type"]
    path = payload["path"]
    if file_type == "pdf":
        return _parse_pdf(
            path,
            payload.get("selected_pages"),
            payload.get("ocr_pages", []),
            payload["ocr_model_root"],
        )
    if file_type == "docx":
        return {
            "text": _parse_docx(path),
            "read_pages": [],
            "scanned_pages": [],
            "ocr_pages": [],
            "total_pages": None,
        }
    if file_type == "html":
        return {
            "text": _parse_html(path),
            "read_pages": [],
            "scanned_pages": [],
            "ocr_pages": [],
            "total_pages": None,
        }
    if file_type == "markdown":
        return {
            "text": _ensure_text_limit(_read_text(path)),
            "read_pages": [],
            "scanned_pages": [],
            "ocr_pages": [],
            "total_pages": None,
        }
    raise AttachmentError("不支持的附件类型。")


def _worker_entry(result_queue, payload: dict[str, Any]) -> None:
    try:
        result_queue.put({"ok": True, "result": _parse_worker(payload)})
    except AttachmentError as exc:
        result_queue.put({"ok": False, "error": str(exc)})
    except Exception:
        result_queue.put({"ok": False, "error": "本地解析失败，文件未发送给模型。"})


class AttachmentStore:
    """Keep pending sources briefly and persist owner-scoped question-answer records."""

    def __init__(self, inbox_root: Path):
        self.root = Path(inbox_root) / "附件问答"
        self.pending_root = self.root / "临时"
        self.record_root = self.root / "记录"
        self.ocr_model_root = self.root / "ocr-models"
        self.lock = threading.RLock()
        for directory in (self.pending_root, self.record_root, self.ocr_model_root):
            directory.mkdir(parents=True, exist_ok=True)
        self.cleanup_expired()

    def _task_dir(self, key: str) -> Path:
        if not TASK_ID_RE.fullmatch(key):
            raise AttachmentError("附件编号无效。")
        target = (self.pending_root / key).resolve()
        if not target.is_relative_to(self.pending_root.resolve()):
            raise AttachmentError("附件编号无效。")
        return target

    def _task_path(self, key: str) -> Path:
        return self._task_dir(key) / "task.json"

    def _source_path(self, key: str) -> Path:
        return self._task_dir(key) / "source.bin"

    def _record_path(self, key: str) -> Path:
        if not RECORD_ID_RE.fullmatch(key):
            raise AttachmentError("附件记录编号无效。")
        return self.record_root / (key + ".json")

    def _task(self, owner: str, key: str) -> dict[str, Any]:
        path = self._task_path(key)
        try:
            task = json.loads(path.read_text("utf-8"))
        except (OSError, ValueError) as exc:
            raise AttachmentError("附件记录已不存在或已损坏，请重新发送文件。") from exc
        if task.get("owner") != digest(owner):
            raise AttachmentError("无权访问该附件。")
        try:
            expired = datetime.fromisoformat(task["expires_at"]) <= now()
        except (KeyError, TypeError, ValueError):
            expired = True
        if expired:
            self._discard(key)
            raise AttachmentError("附件确认已过期，请重新发送文件。")
        return task

    def _save_task(self, key: str, task: dict[str, Any]) -> None:
        atomic_json(self._task_path(key), task)

    def _discard(self, key: str) -> None:
        target = self._task_dir(key)
        if target.exists():
            shutil.rmtree(target)

    def cleanup_expired(self) -> None:
        """Remove only this feature's expired task folders, never adapter media broadly."""
        with self.lock:
            for directory in self.pending_root.iterdir():
                if not directory.is_dir() or not TASK_ID_RE.fullmatch(directory.name):
                    continue
                task_path = directory / "task.json"
                expired = False
                try:
                    task = json.loads(task_path.read_text("utf-8"))
                    expired = datetime.fromisoformat(task["expires_at"]) <= now()
                except (OSError, ValueError, KeyError, TypeError):
                    expired = (
                        directory.stat().st_mtime < time.time() - PENDING_MINUTES * 60
                    )
                if expired:
                    self._discard(directory.name)

    def register(self, owner: str, file_path: str, file_name: str) -> dict[str, Any]:
        """Copy one verified attachment into a short-lived, feature-owned directory."""
        self.cleanup_expired()
        source = Path(file_path)
        if not source.is_file():
            raise AttachmentError("附件暂时不可读取，请重新发送。")
        if source.stat().st_size > MAX_FILE_BYTES:
            raise AttachmentError("单个附件最多 20 MB，请拆分后重新发送。")
        file_type = detect_file_type(str(source), file_name)
        key = uuid.uuid4().hex[:12]
        task_dir = self._task_dir(key)
        task_dir.mkdir(parents=True, exist_ok=False)
        destination = self._source_path(key)
        copied = 0
        try:
            with source.open("rb") as incoming, destination.open("xb") as outgoing:
                while chunk := incoming.read(1024 * 1024):
                    copied += len(chunk)
                    if copied > MAX_FILE_BYTES:
                        raise AttachmentError("附件在读取时超过 20 MB，未处理。")
                    outgoing.write(chunk)
            task = {
                "id": key,
                "owner": digest(owner),
                "file_name": Path(file_name).name or "附件",
                "file_type": file_type,
                "size_bytes": copied,
                "created_at": now().isoformat(),
                "expires_at": (now() + timedelta(minutes=PENDING_MINUTES)).isoformat(),
                "state": "awaiting_question",
                "question": None,
            }
            self._save_task(key, task)
            return task
        except Exception:
            self._discard(key)
            raise

    def set_question(self, owner: str, key: str, question: str) -> dict[str, Any]:
        task = self._task(owner, key)
        question = question.strip()
        if not question or len(question) > 4000:
            raise AttachmentError("请发送 1 至 4000 字的问题。")
        if task.get("state") != "awaiting_question":
            raise AttachmentError("该附件已有关联问题，请确认、取消或重新发送。")
        task.update(question=question, state="awaiting_confirmation")
        self._save_task(key, task)
        return task

    def pending(
        self, owner: str, state: str = "awaiting_question"
    ) -> list[dict[str, Any]]:
        self.cleanup_expired()
        tasks = []
        for directory in self.pending_root.iterdir():
            if not directory.is_dir() or not TASK_ID_RE.fullmatch(directory.name):
                continue
            try:
                task = json.loads((directory / "task.json").read_text("utf-8"))
            except (OSError, ValueError):
                continue
            if task.get("owner") == digest(owner) and task.get("state") == state:
                tasks.append(task)
        return sorted(tasks, key=lambda item: item.get("created_at", ""), reverse=True)

    def confirmation_text(self, task: dict[str, Any]) -> str:
        suffix = (
            "PDF 可附页码，例如：确认读取 " + task["id"] + " 1-3,5"
            if task["file_type"] == "pdf"
            else "确认后将开始本地解析。"
        )
        return (
            f"附件 {task['id']}：{task['file_name']}（{format_bytes(task['size_bytes'])}）\n"
            f"你的问题：{task['question']}\n"
            "将先在本机提取文字，再把文字和问题发送给当前对话模型；"
            "原文件和提取文字只用于本轮，问答会保留供后续查阅，不会自动写入记忆或待整理。\n"
            f"发送“确认读取 {task['id']}”继续，或“取消读取 {task['id']}”删除临时文件。{suffix}"
        )

    def cancel(self, owner: str, key: str) -> None:
        self._task(owner, key)
        self._discard(key)

    def _run_worker(self, payload: dict[str, Any], timeout: int) -> dict[str, Any]:
        context = multiprocessing.get_context("spawn")
        queue = context.Queue(maxsize=1)
        process = context.Process(
            target=_worker_entry, args=(queue, payload), daemon=True
        )
        process.start()
        process.join(timeout)
        if process.is_alive():
            process.terminate()
            process.join(5)
            raise AttachmentError(
                "本地解析超时，文件未发送给模型。请缩小文件或选择更少页码。"
            )
        try:
            result = queue.get(timeout=1)
        except Empty as exc:
            raise AttachmentError("本地解析失败，文件未发送给模型。") from exc
        finally:
            queue.close()
            queue.join_thread()
        if not result.get("ok"):
            raise AttachmentError(
                result.get("error") or "本地解析失败，文件未发送给模型。"
            )
        return result["result"]

    def prepare(
        self,
        owner: str,
        key: str,
        selected_pages: list[int] | None = None,
        ocr_pages: list[int] | None = None,
        *,
        allow_skip_ocr: bool = False,
    ) -> dict[str, Any]:
        """Serialize local parsing while keeping its worker outside the event loop."""
        with self.lock:
            return self._prepare_locked(
                owner,
                key,
                selected_pages,
                ocr_pages,
                allow_skip_ocr=allow_skip_ocr,
            )

    def _prepare_locked(
        self,
        owner: str,
        key: str,
        selected_pages: list[int] | None = None,
        ocr_pages: list[int] | None = None,
        *,
        allow_skip_ocr: bool = False,
    ) -> dict[str, Any]:
        """Parse a confirmed file in a killable worker and return memory-only text."""
        task = self._task(owner, key)
        state = task.get("state")
        if state not in {"awaiting_confirmation", "awaiting_ocr"}:
            raise AttachmentError("该附件当前不能确认读取。")
        if not task.get("question"):
            raise AttachmentError("附件还没有关联问题。")
        ocr_pages = ocr_pages or []
        if task["file_type"] != "pdf" and (selected_pages or ocr_pages):
            raise AttachmentError("只有 PDF 支持页码选择。")
        if task["file_type"] == "pdf" and state == "awaiting_ocr":
            known_scans = set(task.get("scanned_pages", []))
            if not ocr_pages and not allow_skip_ocr:
                raise AttachmentError("请确认需要 OCR 的页码，或发送“略过 OCR 编号”。")
            if not set(ocr_pages).issubset(known_scans):
                raise AttachmentError("OCR 页码必须是提示中的扫描页。")
            selected_pages = task.get("selected_pages") or selected_pages
        result = self._run_worker(
            {
                "path": str(self._source_path(key)),
                "file_type": task["file_type"],
                "selected_pages": selected_pages,
                "ocr_pages": ocr_pages,
                "ocr_model_root": str(self.ocr_model_root),
            },
            120 if ocr_pages else 30,
        )
        if result["scanned_pages"] and not ocr_pages and task["file_type"] == "pdf":
            task.update(
                state="awaiting_ocr",
                scanned_pages=result["scanned_pages"],
                selected_pages=result["read_pages"],
            )
            self._save_task(key, task)
            return {
                "status": "ocr_required",
                "task": task,
                "scanned_pages": result["scanned_pages"],
            }
        task.update(
            state="answering",
            selected_pages=result["read_pages"],
            ocr_pages=result["ocr_pages"],
            prepared_at=now().isoformat(),
        )
        self._save_task(key, task)
        return {
            "status": "ready",
            "task": task,
            "text": result["text"],
            "read_pages": result["read_pages"],
            "ocr_pages": result["ocr_pages"],
        }

    def skip_ocr(self, owner: str, key: str) -> dict[str, Any]:
        task = self._task(owner, key)
        if task.get("state") != "awaiting_ocr":
            raise AttachmentError("该附件当前不在等待 OCR。")
        return self.prepare(
            owner,
            key,
            task.get("selected_pages"),
            [],
            allow_skip_ocr=True,
        )

    def save_answer(
        self,
        owner: str,
        key: str,
        answer: str,
        *,
        conversation_id: str | None,
        umo: str,
        input_tokens: int | None,
        elapsed_ms: int,
        model_calls: int,
    ) -> dict[str, Any] | None:
        """Persist only the permitted question-answer record and remove its source."""
        try:
            task = self._task(owner, key)
        except AttachmentError:
            return None
        if task.get("state") != "answering" or not answer.strip():
            self._discard(key)
            return None
        record = {
            "id": key,
            "owner": digest(owner),
            "file_name": task["file_name"],
            "file_type": task["file_type"],
            "question": task["question"],
            "answer": answer.strip(),
            "created_at": now().isoformat(),
            "read_pages": task.get("selected_pages", []),
            "ocr_pages": task.get("ocr_pages", []),
            "model_calls": max(1, int(model_calls)),
            "input_tokens": input_tokens,
            "elapsed_ms": max(0, int(elapsed_ms)),
            "authority": "附件问答记录，原文未保留；仅代表当时讨论。",
        }
        atomic_json(self._record_path(key), record)
        self._discard(key)
        return record

    def discard_after_response(self, owner: str, key: str) -> None:
        try:
            self._task(owner, key)
        except AttachmentError:
            return
        self._discard(key)

    def list_records(self, owner: str, limit: int = 10) -> list[dict[str, Any]]:
        records = []
        for path in self.record_root.glob("*.json"):
            try:
                record = json.loads(path.read_text("utf-8"))
            except (OSError, ValueError):
                continue
            if record.get("owner") == digest(owner):
                records.append(record)
        return sorted(
            records, key=lambda item: item.get("created_at", ""), reverse=True
        )[:limit]

    def search_records(
        self, owner: str, query: str, limit: int = 5
    ) -> list[dict[str, Any]]:
        terms = search_terms(query)
        if not terms:
            raise AttachmentError("请提供至少两个字的检索词。")
        results = []
        for record in self.list_records(owner, limit=200):
            haystack = "\n".join(
                str(record.get(key, "")) for key in ("file_name", "question", "answer")
            ).lower()
            matched = [term for term in terms if term in haystack]
            if not matched:
                continue
            results.append(
                {
                    "id": record["id"],
                    "file_name": record["file_name"],
                    "question": record["question"][:300],
                    "answer_excerpt": record["answer"][:500],
                    "created_at": record["created_at"],
                    "matched_terms": matched[:8],
                    "authority": record["authority"],
                }
            )
        return results[: max(1, min(8, int(limit)))]

    def read_record(self, owner: str, key: str) -> dict[str, Any]:
        try:
            record = json.loads(self._record_path(key).read_text("utf-8"))
        except (OSError, ValueError) as exc:
            raise AttachmentError("未找到该附件问答记录。") from exc
        if record.get("owner") != digest(owner):
            raise AttachmentError("无权访问该附件问答记录。")
        return {
            key: record.get(key)
            for key in (
                "id",
                "file_name",
                "file_type",
                "question",
                "answer",
                "created_at",
                "read_pages",
                "ocr_pages",
                "authority",
            )
        }

    def delete_record(self, owner: str, key: str) -> dict[str, Any]:
        try:
            record = json.loads(self._record_path(key).read_text("utf-8"))
        except (OSError, ValueError) as exc:
            raise AttachmentError("未找到该附件问答记录。") from exc
        if record.get("owner") != digest(owner):
            raise AttachmentError("无权访问该附件问答记录。")
        self._record_path(key).unlink()
        return record
